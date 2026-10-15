"""工作 → 廠商紀錄：以「案子」（一間公司一個）分開存客戶／廠商資料。

案子 → 分類（例：客戶聯繫單、PU客戶、版本更新，欄位可自訂）→ 一筆筆資料。
可以把 Excel 匯入：每個工作表偵測標題列變成一個分類，同名的分類會合併。
資料只存在主機的資料庫，不會進 git。
"""

import io
import json
import re
from datetime import date, datetime

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from pydantic import BaseModel
from sqlmodel import Session, func, select

from app.database import get_session
from app.models import WorkCase, WorkRow, WorkTable

router = APIRouter(prefix="/api/work", tags=["work"])

MAX_XLSX = 15 * 1024 * 1024


# ---------- 小工具 ----------


def _cols(t: WorkTable) -> list[dict]:
    try:
        cols = json.loads(t.columns or "[]")
        return cols if isinstance(cols, list) else []
    except ValueError:
        return []


def _data(r: WorkRow) -> dict:
    try:
        d = json.loads(r.data or "{}")
        return d if isinstance(d, dict) else {}
    except ValueError:
        return {}


def _next_key(cols: list[dict]) -> str:
    used = {c.get("key") for c in cols}
    n = 1
    while f"c{n}" in used:
        n += 1
    return f"c{n}"


def _clean_columns(raw) -> list[dict]:
    raw = [c for c in raw or [] if isinstance(c, dict)]
    # 先保留既有欄位的 key，新欄位才不會搶到舊欄位的 key（資料會對錯欄）
    reserved = {str(c.get("key") or "").strip() for c in raw} - {""}
    out, seen = [], set()
    for c in raw:
        label = str(c.get("label") or "").strip()
        if not label:
            continue
        key = str(c.get("key") or "").strip()
        if not key or key in seen:
            key = _next_key([{"key": k} for k in reserved | seen])
        seen.add(key)
        out.append({"key": key, "label": label[:60], "masked": bool(c.get("masked"))})
    return out


def _table_out(session, t: WorkTable) -> dict:
    count = session.exec(select(func.count()).select_from(WorkRow).where(WorkRow.table_id == t.id)).one()
    return {"id": t.id, "case_id": t.case_id, "name": t.name, "emoji": t.emoji,
            "columns": _cols(t), "sort": t.sort, "row_count": count}


def _row_out(r: WorkRow) -> dict:
    return {"id": r.id, "table_id": r.table_id, "data": _data(r), "sort": r.sort,
            "updated_at": r.updated_at}


def _get(session, model, id_, what):
    obj = session.get(model, id_)
    if not obj:
        raise HTTPException(404, f"找不到這個{what}")
    return obj


# ---------- 案子 ----------


class CaseIn(BaseModel):
    name: str
    emoji: str = "🏢"
    note: str = ""
    archived: bool = False


@router.get("/cases")
def list_cases(session: Session = Depends(get_session)):
    cases = session.exec(select(WorkCase).order_by(WorkCase.archived, WorkCase.sort, WorkCase.id)).all()
    out = []
    for c in cases:
        tables = session.exec(select(WorkTable).where(WorkTable.case_id == c.id)).all()
        ids = [t.id for t in tables]
        rows = session.exec(select(func.count()).select_from(WorkRow).where(WorkRow.table_id.in_(ids))).one() if ids else 0
        out.append({**c.model_dump(), "table_count": len(tables), "row_count": rows})
    return out


@router.post("/cases", status_code=201)
def create_case(payload: CaseIn, session: Session = Depends(get_session)):
    if not payload.name.strip():
        raise HTTPException(400, "請輸入案子名稱")
    n = session.exec(select(func.count()).select_from(WorkCase)).one()
    c = WorkCase(name=payload.name.strip(), emoji=payload.emoji or "🏢", note=payload.note, sort=n)
    session.add(c)
    session.commit()
    session.refresh(c)
    return c


@router.put("/cases/{case_id}")
def update_case(case_id: int, payload: CaseIn, session: Session = Depends(get_session)):
    c = _get(session, WorkCase, case_id, "案子")
    if not payload.name.strip():
        raise HTTPException(400, "請輸入案子名稱")
    c.name, c.emoji, c.note, c.archived = payload.name.strip(), payload.emoji or "🏢", payload.note, payload.archived
    session.add(c)
    session.commit()
    session.refresh(c)
    return c


@router.delete("/cases/{case_id}", status_code=204)
def delete_case(case_id: int, session: Session = Depends(get_session)):
    c = _get(session, WorkCase, case_id, "案子")
    for t in session.exec(select(WorkTable).where(WorkTable.case_id == c.id)).all():
        for r in session.exec(select(WorkRow).where(WorkRow.table_id == t.id)).all():
            session.delete(r)
        session.delete(t)
    session.delete(c)
    session.commit()


# ---------- 分類 ----------


class TableIn(BaseModel):
    name: str
    emoji: str = "📋"
    columns: list[dict] | None = None


@router.get("/cases/{case_id}/tables")
def list_tables(case_id: int, session: Session = Depends(get_session)):
    _get(session, WorkCase, case_id, "案子")
    tables = session.exec(select(WorkTable).where(WorkTable.case_id == case_id).order_by(WorkTable.sort, WorkTable.id)).all()
    return [_table_out(session, t) for t in tables]


@router.post("/cases/{case_id}/tables", status_code=201)
def create_table(case_id: int, payload: TableIn, session: Session = Depends(get_session)):
    _get(session, WorkCase, case_id, "案子")
    if not payload.name.strip():
        raise HTTPException(400, "請輸入分類名稱")
    cols = _clean_columns(payload.columns) or [{"key": "c1", "label": "名稱", "masked": False},
                                                {"key": "c2", "label": "電話", "masked": False},
                                                {"key": "c3", "label": "聯絡人", "masked": False},
                                                {"key": "c4", "label": "備註", "masked": False}]
    n = session.exec(select(func.count()).select_from(WorkTable).where(WorkTable.case_id == case_id)).one()
    t = WorkTable(case_id=case_id, name=payload.name.strip(), emoji=payload.emoji or "📋",
                  columns=json.dumps(cols, ensure_ascii=False), sort=n)
    session.add(t)
    session.commit()
    session.refresh(t)
    return _table_out(session, t)


@router.put("/tables/{table_id}")
def update_table(table_id: int, payload: TableIn, session: Session = Depends(get_session)):
    t = _get(session, WorkTable, table_id, "分類")
    if not payload.name.strip():
        raise HTTPException(400, "請輸入分類名稱")
    t.name, t.emoji = payload.name.strip(), payload.emoji or "📋"
    if payload.columns is not None:
        cols = _clean_columns(payload.columns)
        if not cols:
            raise HTTPException(400, "至少要有一個欄位")
        t.columns = json.dumps(cols, ensure_ascii=False)
    session.add(t)
    session.commit()
    return _table_out(session, t)


@router.delete("/tables/{table_id}", status_code=204)
def delete_table(table_id: int, session: Session = Depends(get_session)):
    t = _get(session, WorkTable, table_id, "分類")
    for r in session.exec(select(WorkRow).where(WorkRow.table_id == t.id)).all():
        session.delete(r)
    session.delete(t)
    session.commit()


# ---------- 資料 ----------


class RowIn(BaseModel):
    data: dict


def _clean_data(d: dict) -> str:
    return json.dumps({str(k): str(v).strip() for k, v in (d or {}).items() if str(v or "").strip()},
                      ensure_ascii=False)


@router.get("/tables/{table_id}/rows")
def list_rows(table_id: int, session: Session = Depends(get_session)):
    _get(session, WorkTable, table_id, "分類")
    rows = session.exec(select(WorkRow).where(WorkRow.table_id == table_id).order_by(WorkRow.sort, WorkRow.id)).all()
    return [_row_out(r) for r in rows]


@router.post("/tables/{table_id}/rows", status_code=201)
def create_row(table_id: int, payload: RowIn, session: Session = Depends(get_session)):
    _get(session, WorkTable, table_id, "分類")
    last = session.exec(select(func.max(WorkRow.sort)).where(WorkRow.table_id == table_id)).one() or 0
    r = WorkRow(table_id=table_id, data=_clean_data(payload.data), sort=last + 1)
    session.add(r)
    session.commit()
    session.refresh(r)
    return _row_out(r)


@router.put("/rows/{row_id}")
def update_row(row_id: int, payload: RowIn, session: Session = Depends(get_session)):
    r = _get(session, WorkRow, row_id, "資料")
    r.data = _clean_data(payload.data)
    r.updated_at = datetime.utcnow()
    session.add(r)
    session.commit()
    session.refresh(r)
    return _row_out(r)


@router.delete("/rows/{row_id}", status_code=204)
def delete_row(row_id: int, session: Session = Depends(get_session)):
    session.delete(_get(session, WorkRow, row_id, "資料"))
    session.commit()


@router.get("/cases/{case_id}/search")
def search(case_id: int, q: str, session: Session = Depends(get_session)):
    """在這個案子的所有分類裡找（例：找某間公司的電話，不用一個個分類翻）。"""
    words = [w.lower() for w in q.split() if w.strip()]
    if not words:
        return []
    out = []
    for t in session.exec(select(WorkTable).where(WorkTable.case_id == case_id).order_by(WorkTable.sort, WorkTable.id)).all():
        masked = {c["key"] for c in _cols(t) if c.get("masked")}
        for r in session.exec(select(WorkRow).where(WorkRow.table_id == t.id).order_by(WorkRow.sort, WorkRow.id)).all():
            d = _data(r)
            text = " ".join(v for k, v in d.items() if k not in masked).lower()
            if all(w in text for w in words):
                out.append({"table_id": t.id, "table_name": t.name, "table_emoji": t.emoji, "row": _row_out(r)})
                if len(out) >= 200:
                    return out
    return out


# ---------- Excel 匯入 ----------


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d") if v.time() == datetime.min.time() else v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _col_letter(i: int) -> str:
    s = ""
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def _parse_workbook(raw: bytes) -> list[dict]:
    try:
        import openpyxl  # 延遲載入：沒裝也不影響其他功能
    except ImportError:
        raise HTTPException(424, "主機還沒裝 openpyxl：重新執行一次更新（win-restart.bat）就會自動安裝")
    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(400, f"讀不到這個檔案（只支援 .xlsx）：{type(exc).__name__}")

    sheets = []
    for ws in wb.worksheets:
        raw_rows = [list(r) for r in ws.iter_rows(values_only=True)]
        rows = [[_cell(v) for v in r] for r in raw_rows]
        while rows and not any(rows[-1]):
            rows.pop()
            raw_rows.pop()
        if not rows:
            continue
        width = max(len(r) for r in rows)
        rows = [r + [""] * (width - len(r)) for r in rows]

        # 標題列：前 15 列中，第一個「至少 2 格有字、而且都是文字（不是數字）」的列
        header_idx = None
        for i, r in enumerate(raw_rows[:15]):
            vals = [v for v in r if v is not None and str(v).strip()]
            if len(vals) >= 2 and all(isinstance(v, str) and not re.fullmatch(r"[\d\-./: ]+", v.strip()) for v in vals):
                header_idx = i
                break
        # 分類名稱：標題列上方最後一個「只有一格字」的列（例：客戶聯繫單），否則用工作表名稱
        title = ws.title
        for r in rows[: header_idx or 0]:
            filled = [v for v in r if v]
            if len(filled) == 1:
                title = filled[0]

        header = rows[header_idx] if header_idx is not None else [""] * width
        body = [r for r in rows[(header_idx + 1 if header_idx is not None else 0):] if any(r)]
        used = [j for j in range(width) if header[j] or any(r[j] for r in body)]
        labels, seen = [], {}
        for j in used:
            lab = header[j] or f"欄{_col_letter(j)}"
            seen[lab] = seen.get(lab, 0) + 1
            labels.append(lab if seen[lab] == 1 else f"{lab} ({seen[lab]})")
        data = [[r[j] for j in used] for r in body]
        first_fill = sum(1 for r in data if r and r[0]) / len(data) if data else 0
        sheets.append({
            "sheet": ws.title,
            "title": title,
            "header_found": header_idx is not None,
            "columns": labels,
            "rows": data,
            # 第一欄空白的列多半是上一筆的續行（例：同一家公司的第二行聯絡人）
            "merge_default": header_idx is not None and first_fill >= 0.6,
        })
    wb.close()
    return sheets


@router.post("/import/preview")
async def import_preview(file: UploadFile = File(...)):
    raw = await file.read()
    if len(raw) > MAX_XLSX:
        raise HTTPException(400, "檔案太大（上限 15MB）")
    return {"filename": file.filename, "sheets": _parse_workbook(raw)}


class ImportTable(BaseModel):
    name: str
    columns: list[str]
    rows: list[list[str]]


class ImportIn(BaseModel):
    tables: list[ImportTable]


@router.post("/cases/{case_id}/import")
def import_commit(case_id: int, payload: ImportIn, session: Session = Depends(get_session)):
    """寫入匯入的資料：同名分類會合併（欄位依名稱對應，多的欄位自動加上）。"""
    _get(session, WorkCase, case_id, "案子")
    summary = []
    for it in payload.tables:
        name = it.name.strip()
        if not name or not it.columns:
            continue
        t = session.exec(select(WorkTable).where(WorkTable.case_id == case_id, WorkTable.name == name)).first()
        created = t is None
        if created:
            n = session.exec(select(func.count()).select_from(WorkTable).where(WorkTable.case_id == case_id)).one()
            t = WorkTable(case_id=case_id, name=name, columns="[]", sort=n)
            session.add(t)
            session.flush()
        cols = _cols(t)
        by_label = {c["label"]: c["key"] for c in cols}
        keys = []
        for lab in it.columns:
            lab = (lab or "").strip()[:60] or f"欄{len(cols) + 1}"
            if lab not in by_label:
                k = _next_key(cols)
                cols.append({"key": k, "label": lab, "masked": False})
                by_label[lab] = k
            keys.append(by_label[lab])
        t.columns = json.dumps(cols, ensure_ascii=False)
        session.add(t)
        last = session.exec(select(func.max(WorkRow.sort)).where(WorkRow.table_id == t.id)).one() or 0
        added = 0
        for r in it.rows:
            d = {keys[j]: str(v).strip() for j, v in enumerate(r[: len(keys)]) if str(v or "").strip()}
            if d:
                last += 1
                session.add(WorkRow(table_id=t.id, data=json.dumps(d, ensure_ascii=False), sort=last))
                added += 1
        summary.append({"name": name, "created": created, "rows": added})
    session.commit()
    return {"tables": summary}
